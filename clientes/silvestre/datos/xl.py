import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
r = pd.read_pickle('r2.pkl'); tot = pd.read_pickle('tot.pkl'); d = pd.read_pickle('d.pkl')
ORDER = list(tot.index)
F = 'Arial'; H = Font(name=F, bold=True, color='FFFFFF', size=10); N = Font(name=F, size=10); B = Font(name=F, size=10, bold=True)
FILL = PatternFill('solid', fgColor='1E5A23'); BAND = PatternFill('solid', fgColor='EAF9C7'); thin = Side(style='thin', color='E4E4DE')
MONEY = '$ #,##0;-$ #,##0;-'; INT = '#,##0;-#,##0;-'; PCT = '0.0%'
wb = Workbook()

def header(ws, row, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        x = ws.cell(row=row, column=i, value=c); x.font = H; x.fill = FILL; x.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30

NOTA = ('Fuente: export "Jardin_YTD" de Mercado Libre (2026 a la fecha), compartido por el usuario. El export tiene 66.262 publicaciones '
        'de todos los vendedores. Unidades y ventas vienen redondeadas en el export (son estimaciones). Solo se cuentan publicaciones cuya Marca '
        'es la marca oficial; se excluyen "compatible", "alternativo" y listas de varias marcas. Ventas en pesos corrientes.')

# Resumen
ws = wb.active; ws.title = 'Resumen'
ws['A1'] = 'Marcas de Silvestre en Mercado Libre · Jardín 2026'; ws['A1'].font = Font(name=F, bold=True, size=14)
ws['A2'] = NOTA; ws['A2'].font = Font(name=F, size=9, italic=True, color='5F6368'); ws.merge_cells('A2:I2'); ws['A2'].alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[2].height = 42
cols = ['Marca', 'Unidades', 'Ventas $', 'Precio promedio $', '% de las ventas de estas marcas', 'Publicaciones', 'Vendedores', 'Categoría principal', '% unidades en esa categoría']
header(ws, 4, cols, [22, 12, 18, 16, 16, 13, 12, 26, 16])
n = len(ORDER); first, last = 5, 4 + n
for i, b in enumerate(ORDER):
    rr = 5 + i; t = tot.loc[b]
    vals = [b, int(t.unidades), float(t.ventas), f'=IFERROR(C{rr}/B{rr},0)', f'=IFERROR(C{rr}/$C${last+1},0)', int(t.publicaciones), int(t.vendedores), t.categoria_principal, float(t.share_unid_cat_ppal)]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=rr, column=j, value=v); c.font = N; c.border = Border(bottom=thin)
    for j, f in [(2, INT), (3, MONEY), (4, MONEY), (5, PCT), (6, INT), (7, INT), (9, PCT)]: ws.cell(row=rr, column=j).number_format = f
tr = last + 1
for j, v in enumerate(['Total', f'=SUM(B{first}:B{last})', f'=SUM(C{first}:C{last})', f'=IFERROR(C{tr}/B{tr},0)', f'=SUM(E{first}:E{last})', f'=SUM(F{first}:F{last})'], 1):
    c = ws.cell(row=tr, column=j, value=v); c.font = B; c.fill = BAND
for j, f in [(2, INT), (3, MONEY), (4, MONEY), (5, PCT), (6, INT)]: ws.cell(row=tr, column=j).number_format = f
ws.cell(row=tr + 2, column=1, value='Picasso (semillas de césped) no aparece en el export: las semillas no están en la categoría Jardín que se exportó.').font = Font(name=F, size=9, italic=True, color='5F6368')
ws.cell(row=tr + 3, column=1, value='MTD / Yard Machines casi no tiene ventas de equipos: lo que aparece son repuestos.').font = Font(name=F, size=9, italic=True, color='5F6368')
ws.freeze_panes = 'B5'

# Top por marca
ws2 = wb.create_sheet('Top por marca')
ws2['A1'] = 'Lo más vendido de cada marca · top 5 máquinas y herramientas + top 3 repuestos y accesorios'; ws2['A1'].font = Font(name=F, bold=True, size=13)
cols2 = ['Marca', 'Puesto', 'Tipo', 'Modelo', 'Publicación más vendida', 'Categoría', 'Unidades', 'Ventas $', 'Precio promedio $', 'Vendedores', 'Publicaciones']
header(ws2, 3, cols2, [20, 8, 20, 18, 62, 26, 11, 16, 15, 11, 13])
rr = 4
for b in ORDER:
    g = r[r.Marca == b]
    for tipo, k in [('Máquina / herramienta', 5), ('Repuesto / accesorio', 3)]:
        x = g[g.Tipo == tipo].head(k)
        for pos, (_, row) in enumerate(x.iterrows(), 1):
            vals = [b, pos, tipo, row.Modelo, row['Publicación más vendida'], row.Categoría, int(row.Unidades), float(row['Ventas $']), f'=IFERROR(H{rr}/G{rr},0)', int(row.Vendedores), int(row.Publicaciones)]
            for j, v in enumerate(vals, 1):
                c = ws2.cell(row=rr, column=j, value=v); c.font = B if (j == 1 and pos == 1 and tipo.startswith('M')) else N; c.border = Border(bottom=thin)
                if ORDER.index(b) % 2 == 0: c.fill = BAND
            for j, f in [(7, INT), (8, MONEY), (9, MONEY), (10, INT), (11, INT)]: ws2.cell(row=rr, column=j).number_format = f
            rr += 1
ws2.freeze_panes = 'A4'; ws2.auto_filter.ref = f'A3:K{rr-1}'

# Todos los productos
ws3 = wb.create_sheet('Todos los productos')
cols3 = ['Marca', 'Tipo', 'Categoría', 'Subcategoría', 'Modelo', 'Publicación más vendida', 'Unidades', 'Ventas $', 'Precio promedio $', 'Publicaciones', 'Vendedores']
header(ws3, 1, cols3, [20, 20, 26, 18, 18, 62, 11, 16, 15, 13, 11])
rs = r.sort_values(['Marca', 'Unidades'], ascending=[True, False])
for i, (_, row) in enumerate(rs.iterrows(), 2):
    vals = [row.Marca, row.Tipo, row.Categoría, row.Subcategoría, row.Modelo, row['Publicación más vendida'], int(row.Unidades), float(row['Ventas $']), f'=IFERROR(H{i}/G{i},0)', int(row.Publicaciones), int(row.Vendedores)]
    for j, v in enumerate(vals, 1): ws3.cell(row=i, column=j, value=v).font = N
    for j, f in [(7, INT), (8, MONEY), (9, MONEY)]: ws3.cell(row=i, column=j).number_format = f
ws3.freeze_panes = 'A2'; ws3.auto_filter.ref = f'A1:K{len(rs)+1}'
wb.save('silvestre-marcas-mas-vendidas-meli-2026.xlsx'); print('rows top', rr - 4, 'all', len(rs))
